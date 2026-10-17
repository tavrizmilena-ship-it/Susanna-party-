"""
Локальная имитация потока Power Automate — ТОЛЬКО для тестов на своём компьютере.
Реализует тот же контракт, что и настоящий поток (см. docs/EXCEL_SETUP.md):

  POST /flow   заголовок x-rsvp-secret, JSON {timestamp, guestName, willAttend, totalPeople,
               additionalGuests, submissionId}
  → 200 {"status": "saved"} | {"status": "duplicate"};  401 при неверном секрете.

Пишет строки в настоящий .xlsx (в таблицу RSVP), запросы обрабатываются по одному,
как в потоке с ограничением параллелизма = 1.

Запуск: python3 scripts/mock_flow.py путь/к/копии.xlsx [порт] [секрет]
"""
import json
import sys
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

from openpyxl import load_workbook

PATH = sys.argv[1]
PORT = int(sys.argv[2]) if len(sys.argv) > 2 else 7071
SECRET = sys.argv[3] if len(sys.argv) > 3 else "local-test-secret"
LOCK = threading.Lock()
FIELDS = ["timestamp", "guestName", "willAttend", "totalPeople", "additionalGuests", "submissionId"]


def append_row(data):
    wb = load_workbook(PATH)
    ws = wb["RSVP Responses"]
    table = ws.tables["RSVP"]
    first, last = table.ref.split(":")
    last_row = int("".join(ch for ch in last if ch.isdigit()))
    ids = {ws.cell(row=r, column=6).value for r in range(2, last_row + 1)}
    if data["submissionId"] in ids:
        return "duplicate"
    # первая пустая строка таблицы используется, иначе таблица расширяется
    target = 2 if all(ws.cell(row=2, column=c).value in (None, "") for c in range(1, 7)) else last_row + 1
    for c, key in enumerate(FIELDS, start=1):
        ws.cell(row=target, column=c, value=data[key])
    table.ref = f"A1:F{max(target, last_row)}"
    wb.save(PATH)
    return "saved"


class Handler(BaseHTTPRequestHandler):
    def do_POST(self):
        if self.path != "/flow":
            return self._send(404, {"error": "not found"})
        if self.headers.get("x-rsvp-secret") != SECRET:
            return self._send(401, {"error": "unauthorized"})
        body = json.loads(self.rfile.read(int(self.headers.get("content-length", 0))) or b"{}")
        if any(k not in body for k in FIELDS):
            return self._send(400, {"error": "bad schema"})
        with LOCK:
            status = append_row(body)
        self._send(200, {"status": status})

    def _send(self, code, obj):
        raw = json.dumps(obj).encode()
        self.send_response(code)
        self.send_header("content-type", "application/json")
        self.send_header("content-length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)

    def log_message(self, *args):
        pass


print(f"mock flow on http://127.0.0.1:{PORT}/flow -> {PATH}", flush=True)
ThreadingHTTPServer(("127.0.0.1", PORT), Handler).serve_forever()
