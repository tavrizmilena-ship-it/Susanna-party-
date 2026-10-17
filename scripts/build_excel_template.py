"""
Создаёт пустой шаблон Susanna_Birthday_Party_RSVP.xlsx:
  - лист «RSVP Responses» с таблицей Excel «RSVP» (сюда сайт добавляет строки);
  - лист «Summary» с формулами, которые пересчитываются автоматически;
  - лист «Invited» — необязательный список приглашённых для подсчёта «ещё не ответили».

Запуск:  python3 scripts/build_excel_template.py [путь_к_файлу]
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "excel/Susanna_Birthday_Party_RSVP.xlsx"

FOREST = "35432F"
IVORY = "F5F0E6"
PAPER = "E9E0D2"
PINK = "CBA7A0"
MAX_ROW = 2000

COLUMNS = [
    ("Timestamp", 21),
    ("Guest name", 30),
    ("Will attend?", 14),
    ("Total people", 14),
    ("Additional guests", 18),
    ("Submission ID", 38),
]

wb = Workbook()

# ---------------------------- RSVP Responses ----------------------------
ws = wb.active
ws.title = "RSVP Responses"
for i, (name, width) in enumerate(COLUMNS, start=1):
    c = ws.cell(row=1, column=i, value=name)
    c.font = Font(bold=True, color=IVORY, name="Calibri", size=11)
    c.fill = PatternFill("solid", fgColor=FOREST)
    c.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
    ws.column_dimensions[c.column_letter].width = width
ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"

# В таблице Excel должна быть хотя бы одна строка данных — оставляем её пустой.
table = Table(displayName="RSVP", ref="A1:F2")
table.tableStyleInfo = TableStyleInfo(name="TableStyleLight15", showRowStripes=True)
ws.add_table(table)
# Служебный столбец можно скрыть — он нужен для защиты от дублей.
ws.column_dimensions["F"].hidden = False

# -------------------------------- Summary --------------------------------
s = wb.create_sheet("Summary")
s.sheet_view.showGridLines = False
s.column_dimensions["A"].width = 46
s.column_dimensions["B"].width = 18
s.column_dimensions["C"].width = 60

s["A1"] = "Susánna’s birthday party — сводка ответов"
s["A1"].font = Font(bold=True, size=16, color=FOREST)
s["A2"] = "Обновляется автоматически при каждом новом ответе."
s["A2"].font = Font(italic=True, size=10, color="5B6250")

R = f"'RSVP Responses'!$B$2:$B${MAX_ROW}"
A = f"'RSVP Responses'!$C$2:$C${MAX_ROW}"
T = f"'RSVP Responses'!$D$2:$D${MAX_ROW}"
INV = f"Invited!$A$2:$A${MAX_ROW}"

rows = [
    ("Получено ответов", f"=COUNTA({R})", "Все анкеты, включая повторные от одного гостя."),
    ("Подтвердили присутствие", f'=COUNTIF({A},"Да")', ""),
    ("Не смогут прийти", f'=COUNTIF({A},"Нет")', ""),
    ("Ожидаемое число гостей", f"=SUM({T})", "Сумма «Total people» — включая самих приглашённых."),
    ("Приглашено по списку", f"=COUNTA({INV})", "Заполните лист «Invited», чтобы видеть, кто ещё не ответил."),
    (
        "Ещё не ответили",
        f'=IF(B9=0,"—",SUMPRODUCT(({INV}<>"")*(COUNTIF({R},{INV})=0)))',
        "Сравнение по имени: пишите имена в «Invited» так же, как гости в анкете.",
    ),
]
thin = Side(style="thin", color=PAPER)
s["A4"], s["B4"], s["C4"] = "Показатель", "Значение", "Комментарий"
for c in ("A4", "B4", "C4"):
    s[c].font = Font(bold=True, color=IVORY)
    s[c].fill = PatternFill("solid", fgColor=FOREST)
for i, (label, formula, note) in enumerate(rows, start=5):
    s[f"A{i}"] = label
    s[f"B{i}"] = formula
    s[f"C{i}"] = note
    s[f"A{i}"].font = Font(size=12, color=FOREST)
    s[f"B{i}"].font = Font(size=14, bold=True, color=FOREST)
    s[f"B{i}"].alignment = Alignment(horizontal="center")
    s[f"C{i}"].font = Font(size=9, italic=True, color="5B6250")
    s[f"C{i}"].alignment = Alignment(wrap_text=True, vertical="center")
    for col in "ABC":
        s[f"{col}{i}"].border = Border(bottom=thin)
        s[f"{col}{i}"].fill = PatternFill("solid", fgColor=IVORY)
    s.row_dimensions[i].height = 28

s["A13"] = "Если гость отправил анкету повторно, учитываются обе строки — удалите устаревшую на листе «RSVP Responses»."
s["A13"].font = Font(size=9, italic=True, color="5B6250")

# -------------------------------- Invited --------------------------------
inv = wb.create_sheet("Invited")
inv["A1"], inv["B1"] = "Guest name", "Responded?"
for c in ("A1", "B1"):
    inv[c].font = Font(bold=True, color=IVORY)
    inv[c].fill = PatternFill("solid", fgColor=PINK)
inv.column_dimensions["A"].width = 34
inv.column_dimensions["B"].width = 16
inv["B2"] = f'=IF(A2="","",IF(COUNTIF({R},A2)>0,"Да","Нет"))'
inv_table = Table(displayName="Invited", ref="A1:B2")
inv_table.tableStyleInfo = TableStyleInfo(name="TableStyleLight10", showRowStripes=True)
inv.add_table(inv_table)
inv["D1"] = "Впишите сюда имена приглашённых (по одному в строке) — сводка покажет, кто ещё не ответил."
inv["D1"].font = Font(italic=True, size=9, color="5B6250")

dv = DataValidation(type="list", formula1='"Да,Нет"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"C2:C{MAX_ROW}")

wb.active = 0
wb.save(OUT)
print("saved", OUT)
